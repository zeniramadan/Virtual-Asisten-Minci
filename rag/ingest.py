import os
import sys
import time
import hashlib
import shutil

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import chromadb
import ollama
import docx as docx_lib
import pdfplumber

from llama_index.core import Document
from llama_index.core.node_parser import SentenceSplitter

import rag.config as config
import ollama_utils

SUPPORTED_EXTS = (".docx", ".pdf")

def _line(char="─", n=70):
    print(char * n)

def embed_text(text: str):
    resp = ollama.embeddings(model=config.EMBED_MODEL, prompt=text)
    return resp["embedding"]

def build_chunk_id(file_name: str, idx: int, text: str) -> str:
    h = hashlib.md5(text.encode("utf-8")).hexdigest()[:8]
    safe_name = os.path.splitext(os.path.basename(file_name))[0]
    return f"{safe_name}-{idx}-{h}"

def _get_doc_title(file_name: str) -> str:
    """
    Ambil judul deskriptif dokumen dari config.DOCUMENT_TITLES (mis. "PMB" ->
    "Penerimaan Mahasiswa Baru (PMB)"). Judul ini ditempel di depan teks SAAT
    di-embed (bukan di teks yang disimpan/ditampilkan), supaya embedding tiap
    chunk "tahu" ia berasal dari kategori dokumen apa -> mengurangi salah
    tarik antar-dokumen yang kata-katanya mirip (mis. "pendaftaran" vs
    "registrasi").
    """
    stem = os.path.splitext(os.path.basename(file_name))[0]
    return getattr(config, "DOCUMENT_TITLES", {}).get(stem, stem)

def _looks_like_heading(line: str) -> bool:
    """
    Heuristik deteksi judul section (mis. "PERSYARATAN ADMINISTRASI",
    "SOP REGISTRASI DAN HERREGISTRASI"): baris pendek, 2-10 kata, dan
    semua huruf di dalamnya kapital.
    """
    line = line.strip()
    if not line or len(line) > 80:
        return False
    words = line.split()
    if not (2 <= len(words) <= 10):
        return False
    letters = [ch for ch in line if ch.isalpha()]
    if not letters:
        return False
    return all(ch.isupper() for ch in letters)

def _split_into_sections(text: str) -> list:
    """
    Pecah teks satu dokumen jadi beberapa section berdasarkan baris heading
    (ALL CAPS). Tujuannya supaya nanti saat di-chunk, SentenceSplitter tidak
    menggabungkan ekor satu section dengan awal section lain yang topiknya
    beda (mis. akhir "Cetak KRS" nyambung ke awal "SOP Registrasi") ->
    karena SentenceSplitter memproses tiap Document secara terpisah, dan di
    sini tiap section jadi Document sendiri-sendiri.
    """
    lines = text.split("\n")
    sections = []
    current_lines = []
    for line in lines:
        if _looks_like_heading(line) and current_lines:
            sections.append("\n".join(current_lines).strip())
            current_lines = [line]
        else:
            current_lines.append(line)
    if current_lines:
        sections.append("\n".join(current_lines).strip())
    return [s for s in sections if s.strip()]

def load_documents():
    if not os.path.isdir(config.DOCUMENTS_DIR) or not os.listdir(config.DOCUMENTS_DIR):
        print(f"[!] Folder dokumen kosong: {config.DOCUMENTS_DIR}")
        sys.exit(1)

    file_paths = list(_iter_source_files())
    if not file_paths:
        print(f"[!] Tidak ada file .docx / .pdf di: {config.DOCUMENTS_DIR}")
        sys.exit(1)

    documents = []
    for path in file_paths:
        ext = os.path.splitext(path)[1].lower()
        file_name = os.path.basename(path)
        try:
            if ext == ".docx":
                text = _extract_docx_text(path)
            elif ext == ".pdf":
                text = _extract_pdf_text(path)
            else:
                continue
        except Exception as e:
            print(f"[!] Gagal membaca {file_name}: {e}")
            continue

        if not text.strip():
            print(f"[!] {file_name}: tidak ada teks yang terbaca (PDF hasil scan?) -> dilewati")
            continue

        sections = _split_into_sections(text) or [text]
        for sec_idx, sec_text in enumerate(sections, start=1):
            documents.append(Document(text=sec_text, metadata={"file_name": file_name, "section_index": sec_idx}))
    return documents

def _iter_source_files():
    for dirpath, _, filenames in os.walk(config.DOCUMENTS_DIR):
        for fname in sorted(filenames):
            if fname.lower().endswith(SUPPORTED_EXTS):
                yield os.path.join(dirpath, fname)

def _table_to_markdown(rows) -> str:
    cleaned = [[(c or "").strip() for c in row] for row in rows]
    cleaned = [r for r in cleaned if any(r)]
    if not cleaned: return ""
    header = cleaned[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for row in cleaned[1:]:
        row = (row + [""] * len(header))[:len(header)]
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)

def _extract_docx_text(path: str) -> str:
    doc = docx_lib.Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    body_text = "\n".join(paragraphs)
    table_blocks = []
    for i, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        md = _table_to_markdown(rows)
        if md: table_blocks.append(f"[Tabel {i}]\n{md}")
    return "\n\n".join(part for part in [body_text, *table_blocks] if part)

def _extract_pdf_text(path: str) -> str:
    parts = []
    with pdfplumber.open(path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text: parts.append(f"[Halaman {page_num}]\n{text}")
            for t_idx, rows in enumerate(page.extract_tables(), start=1):
                md = _table_to_markdown(rows)
                if md: parts.append(f"[Halaman {page_num} - Tabel {t_idx}]\n{md}")
    return "\n\n".join(parts)

def chunk_documents(documents):
    parser = SentenceSplitter(chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP)
    return parser.get_nodes_from_documents(documents)

def main():
    start = time.time()
    _line("=")
    print("MINCI RAG — INGEST DOKUMEN")

    ollama_utils.ensure_ready([config.EMBED_MODEL])

    documents = load_documents()
    if not documents:
        print("[!] Tidak ada dokumen yang berhasil dibaca -> ingest dibatalkan.")
        sys.exit(1)
    nodes = chunk_documents(documents)

    if os.path.isdir(config.CHROMA_DIR):
        try:
            shutil.rmtree(config.CHROMA_DIR)
        except PermissionError:
            print(f"[!] Folder database tidak bisa dihapus: {config.CHROMA_DIR}")
            print("    Hentikan dulu server / proses lain yang sedang memakai database (run_server.py, query.py), lalu coba lagi.")
            sys.exit(1)
    os.makedirs(config.CHROMA_DIR, exist_ok=True)

    client = chromadb.PersistentClient(path=config.CHROMA_DIR)
    collection = client.create_collection(name=config.COLLECTION_NAME, metadata={"hnsw:space": "cosine"})

    ids, texts, embeddings, metadatas = [], [], [], []
    for i, node in enumerate(nodes, start=1):
        text = node.get_content().strip()
        if not text: continue
        file_name = node.metadata.get("file_name", "unknown")
        chunk_id = build_chunk_id(file_name, i, text)

        # Contextual embedding: teks yang DISIMPAN (`texts`) tetap teks asli,
        # tapi teks yang DI-EMBED ditambah judul dokumen di depan, supaya
        # vektor embedding-nya membawa konteks kategori dokumen.
        doc_title = _get_doc_title(file_name)
        embed_input = f"Dokumen: {doc_title}\n\n{text}"
        embedding = embed_text(embed_input)

        ids.append(chunk_id)
        texts.append(text)
        embeddings.append(embedding)
        metadatas.append({"file_name": file_name, "chunk_index": i})
        print(f"  [{i:>4}/{len(nodes)}] {file_name:<30} | {len(text):>4} chars")

    if ids:
        collection.add(ids=ids, documents=texts, embeddings=embeddings, metadatas=metadatas)

    elapsed = time.time() - start
    print(f"SELESAI. {len(ids)} chunk tersimpan dalam {elapsed:.1f} detik.")

if __name__ == "__main__":
    main()
